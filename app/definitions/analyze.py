"""Reprodukovatelná extrakce dodaných podkladů; nevykonává metodická pravidla."""
import argparse
from datetime import date
import json
from pathlib import Path
import re
import subprocess

from openpyxl import load_workbook

from app.definitions.catalog import register, verify
from app.definitions.intake import digest
from app.definitions.workbooks import cell_value, record_changes, table_records
from app.xml import parse_local_xml
from app.xml.bundle import BundleValidator

XS = {"xs": "http://www.w3.org/2001/XMLSchema"}


def schema_elements(directory: Path, root: Path) -> list[dict]:
    elements = []
    for path in sorted(directory.glob("*.xsd")):
        schema = parse_local_xml(path.read_bytes())
        for node in schema.xpath(".//xs:element", namespaces=XS):
            documentation = " ".join(node.xpath("./xs:annotation/xs:documentation//text()", namespaces=XS)).strip()
            elements.append({"file": path.relative_to(root).as_posix(), "line": node.sourceline,
                             "name": node.get("name"), "ref": node.get("ref"), "type": node.get("type"),
                             "minOccurs": node.get("minOccurs", "1"), "maxOccurs": node.get("maxOccurs", "1"),
                             "namespace": schema.get("targetNamespace"),
                             "attribute_ids": sorted(set(re.findall(r"\bID\s*(\d{5})\b", documentation))),
                             "documentation": documentation})
    return elements


def write(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    options = {"separators": (",", ":")} if path.parent.name == "analysis" else {"indent": 2}
    path.write_text(json.dumps(value, ensure_ascii=False, sort_keys=True, **options) + "\n", encoding="utf-8")


def analyze(root: Path) -> dict:
    root = root.resolve()
    output = root / "analysis"
    output.mkdir(exist_ok=True)
    intake = json.loads((root / "intake.json").read_text())
    extracted = json.loads((root / "extraction.json").read_text())
    inventory, tables, bundles = [], {}, []
    for item in intake["documents"]:
        path = root / item["file"]
        if digest(path) != item["sha256"]:
            raise ValueError("Originál neodpovídá inventuře.")
        cohort = path.parent.name  # grouping only, NEVER an effective date
        record = dict(item)
        record.update({"effective_from": None, "effective_until": None, "effective_status": "unconfirmed",
                       "version_evidence": [], "publication_date": None})
        if path.suffix == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                if "Verze" in workbook:
                    history = [(i, row) for i, row in enumerate(workbook["Verze"].values, 1) if i > 1 and row[0] is not None]
                    number, last = history[-1]
                    record["version"] = str(last[0])
                    record["publication_date"] = str(last[1])
                    record["version_evidence"] = [f"Verze!A{number}: {last[0]}", f"Verze!B{number}: {last[1]} (publikace, nikoliv účinnost)"]
                else:
                    rows = list(workbook["Titulní strana"].values)
                    row = next((i, r) for i, r in enumerate(rows, 1) if "Verze dokumentu" in str(r[0]))
                    record["version"] = str(row[1][1])
                    record["version_evidence"] = [f"Titulní strana!B{row[0]}: {record['version']}"]
                record["worksheets"] = {s.title: {"rows": s.max_row, "columns": s.max_column} for s in workbook}
                if "slovnik" in path.name:
                    kind, sheet, header, family = "data_dictionary", "SLOVNÍK", 2, "dictionary"
                    codes = {}
                    for s in workbook:
                        if s.title.startswith("CIS "):
                            codes[s.title] = [[cell_value(v) for v in r] for r in s.values]
                    write(output / f"codelists-{cohort}.json", {"source": item["file"], "source_sha256": item["sha256"], "sheets": codes})
                elif "katalog" in path.name:
                    kind, sheet, header, family = "control_catalog", "MH", 1, "controls"
                    parameters = [[cell_value(v) for v in r] for r in workbook["Parametrické konstanty"].values]
                    write(output / f"parameters-{cohort}.json", {"source": item["file"], "source_sha256": item["sha256"], "rows": parameters})
                else:
                    kind, sheet, header, family = "scenarios", "Interakce", 3, "interactions"
            finally:
                workbook.close()
            records = table_records(path.read_bytes(), sheet, header)
            tables[(family, cohort)] = records
            write(output / f"{family}-{cohort}.json", {"source": item["file"], "source_sha256": item["sha256"], "records": records})
            record["kind"] = kind
        elif path.suffix == ".pdf":
            text = subprocess.run(["pdftotext", "-layout", str(path), "-"], check=True, capture_output=True).stdout.decode("utf-8")
            (output / f"{cohort}-{path.stem}.txt").write_text(text, encoding="utf-8")
            versions = re.findall(r"Verze\s+(\d+(?:\.\d+)+)", text[:1800])
            if not versions:
                raise ValueError("Verzi PDF nelze doložit titulní stranou.")
            record["version"] = versions[-1] if "revize" in path.name else versions[0]
            record["version_evidence"] = ["PDF strana 1: " + text[:500].strip()]
            match = re.search(r"Platné pro období od\s+(\d+)\.(\d+)\.(\d+)", text[:1000])
            if match:
                day, month, year = map(int, match.groups())
                record["effective_from"] = date(year, month, day).isoformat()
                record["effective_status"] = "declared_in_document_not_legally_reviewed"
                record["effective_evidence"] = "PDF strana 1: " + match.group(0)
            if "revize" in path.name:
                record["warning"] = "Titulní strana revize obsahuje překryv Verze 1.4.13Verze 1.4.14; nejde o samostatnou novou verzi."
            record["kind"] = "methodology"
        elif path.suffix == ".zip":
            archive = next(a for a in extracted["archives"] if a["source_file"] == item["file"] and not a["chain"])
            if "protokoly" in path.name:
                record["version"] = "1.5"
                record["version_evidence"] = ["Vnitřní adresář a DOCX: MPSV-JMHZ-protokoly-MH v1.5; nejde o vstupní měsíční hlášení."]
            else:
                candidates = [p for p in extracted["published"] if p["source_file"] == item["file"] and p["file"].endswith("externi_xsd/jmhzPodani.xsd")]
                entry = root / candidates[0]["file"]
                xsd = parse_local_xml(entry.read_bytes())
                text = " ".join(xsd.xpath(".//xs:documentation/text()", namespaces={"xs": "http://www.w3.org/2001/XMLSchema"}))
                versions = re.findall(r"Verze\s+(\d+(?:\.\d+)+)", text)
                record["version"] = versions[-1] if versions else xsd.get("version")
                record["version_evidence"] = ["jmhzPodani.xsd/@version=" + xsd.get("version"), "Poslední verze v dokumentaci XSD=" + record["version"]]
                record["xsd_attribute_version"] = xsd.get("version")
                validator = BundleValidator(entry, entry.parent)
                bundles.append({"id": f"jmhz-{record['version']}", "cohort_label": cohort,
                                "bundle_release": record["version"], "xsd_attribute_version": xsd.get("version"),
                                "entry": entry.relative_to(root).as_posix(), "root": entry.parent.relative_to(root).as_posix(),
                                "namespace": xsd.get("targetNamespace"), "compiled_dependencies": validator.resolver.loaded,
                                "origin": "unverified", "effective_from": None, "effective_until": None,
                                "selection_status": "explicit_selection_only_effectivity_not_inferred"})
            record["kind"] = "archive"
            record["archive_members"] = len(archive["members"])
        else:
            raise ValueError("Neznámý dokument v inventuře.")
        inventory.append(record)
        register(root, path, record["source_url"], record["version"], record["kind"], record["effective_from"], record["effective_until"])

    originals = {r["file"]: r for r in inventory}
    for member in extracted["published"]:
        path = root / member["file"]
        parent = originals[member["source_file"]]
        version = parse_local_xml(path.read_bytes()).get("version") if path.suffix == ".xsd" else parent["version"]
        register(root, path, member["source_url"], version or parent["version"], "xsd" if path.suffix == ".xsd" else "archive")

    manifest = json.loads((root / "manifest.json").read_text())
    members = {m["file"]: m for m in extracted["published"]}
    for record in manifest["documents"]:
        source = originals.get(record["file"])
        if source:
            record["source_url_basis"] = source["source_url_basis"]
            record["original_path"] = source["original_path"]
            record["version_evidence"] = source["version_evidence"]
            record["effective_status"] = source["effective_status"]
            record["original_downloaded_at"] = None
            if "effective_evidence" in source:
                record["effective_evidence"] = source["effective_evidence"]
        elif record["file"] in members:
            source = members[record["file"]]
            record["container_chain"] = source["container_chain"]
            record["source_file"] = source["source_file"]
            record["bundle_release"] = originals[source["source_file"]]["version"]
            record["effective_status"] = "unconfirmed"
    write(root / "manifest.json", manifest)
    write(root / "bundles.json", {"bundles": bundles})
    write(output / "inventory.json", {"documents": inventory})
    differences = {family: record_changes(tables[(family, "2026")], tables[(family, "2027")]) for family in ("dictionary", "controls", "interactions")}
    codes = {year: json.loads((output / f"codelists-{year}.json").read_text())["sheets"] for year in ("2026", "2027")}
    differences["codelists"] = {"added_sheets": sorted(codes["2027"].keys() - codes["2026"].keys()),
                                  "removed_sheets": sorted(codes["2026"].keys() - codes["2027"].keys()),
                                  "changed_sheets": sorted(k for k in codes["2026"].keys() & codes["2027"].keys() if codes["2026"][k] != codes["2027"][k])}
    write(output / "differences.json", differences)
    indexes = {}
    for bundle in bundles:
        cohort = bundle["cohort_label"]
        elements = schema_elements(root / bundle["root"], root)
        indexes[cohort] = elements
        write(output / f"xsd-elements-{cohort}.json", elements)
        controls = tables[("controls", cohort)]
        candidates = {}
        for id_, attribute in tables[("dictionary", cohort)].items():
            fields = attribute["fields"]
            value = lambda name: fields.get(name, {}).get("value")
            if not value("MĚSÍČNÍ"):
                continue
            mapping_name = next((k for k in fields if k.startswith("XSD MAPOVÁNÍ MH")), None)
            related = []
            for control_id, control in controls.items():
                refs = next((v["value"] for k, v in control["fields"].items() if k.startswith("Seznam ID atributů")), None)
                if id_ in re.findall(r"\b\d{5}\b", str(refs)):
                    related.append(control_id)
            candidates[id_] = {"name": value("NÁZEV ATRIBUTU"), "type": value("DATOVÝ TYP"),
                                "source_dictionary": next(d["file"] for d in inventory if d["kind"] == "data_dictionary" and d["file"].startswith(cohort + "/")),
                                "source_control_catalog": next(d["file"] for d in inventory if d["kind"] == "control_catalog" and d["file"].startswith(cohort + "/")),
                                "dictionary_row": attribute["row"], "dictionary_path_raw": fields.get(mapping_name, {}),
                                "codelist_reference": fields.get("ČÍSELNÍK"), "related_control_ids": related,
                                "xsd_declarations": [e for e in elements if id_ in e["attribute_ids"]],
                                "status": "working_evidence_not_executable_mapping"}
        write(output / f"mapping-evidence-{cohort}.json", candidates)
    xsd_delta = {}
    for filename in sorted({Path(e["file"]).name for es in indexes.values() for e in es}):
        def signatures(year: str) -> dict:
            result = {}
            for e in indexes[year]:
                if Path(e["file"]).name != filename:
                    continue
                key = e["name"] or e["ref"]
                result.setdefault(key, []).append({k: e[k] for k in ("type", "minOccurs", "maxOccurs", "attribute_ids", "namespace")})
            return result
        left, right = signatures("2026"), signatures("2027")
        xsd_delta[filename] = {"added_names": sorted(right.keys() - left.keys()), "removed_names": sorted(left.keys() - right.keys()),
                               "changed_declarations": {k: {"before": left[k], "after": right[k]} for k in left.keys() & right.keys() if left[k] != right[k]}}
    write(output / "xsd-differences.json", xsd_delta)
    failures = verify(root)
    if failures:
        raise ValueError("Ověření manifestu selhalo.")
    return {"documents": len(inventory), "registered": len(manifest["documents"]), "bundles": len(bundles),
            "differences": {k: {field: len(v) for field, v in value.items()} for k, value in differences.items()}}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path("schemas/official"))
    print(json.dumps(analyze(parser.parse_args().root), ensure_ascii=False, indent=2))
