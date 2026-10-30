"""Čtení technických XLSX jako dokumentů, nikdy vyhodnocování vzorců."""
from datetime import date, datetime
from io import BytesIO
import hashlib
import json

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.definitions.archives import zip_inventory


def cell_value(value: object) -> object:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def workbook_snapshot(content: bytes) -> dict:
    zip_inventory(content)
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    result, count = {}, 0
    try:
        for sheet in workbook:
            digest, nonempty, formulas = hashlib.sha256(), 0, 0
            if (sheet.max_row or 0) * (sheet.max_column or 0) > 1_000_000:
                raise ValueError("XLSX překračuje limit buněk listu.")
            for row in sheet:
                for cell in row:
                    count += 1
                    if count > 3_000_000:
                        raise ValueError("XLSX překračuje limit buněk sešitu.")
                    if cell.value is not None:
                        nonempty += 1
                        formulas += int(cell.data_type == "f")
                        digest.update(json.dumps([cell.coordinate, cell.data_type, cell_value(cell.value)], ensure_ascii=False, sort_keys=True).encode())
            result[sheet.title] = {"rows": sheet.max_row, "columns": sheet.max_column, "nonempty": nonempty,
                                   "formulas_retained_not_executed": formulas, "cells_sha256": digest.hexdigest()}
        return {"sheets": result, "interpretation": "not_validated"}
    finally:
        workbook.close()


def table_records(content: bytes, sheet_name: str, header_row: int, key_column: int = 1) -> dict[str, dict]:
    """Preserve source sheet/row/column; duplicate IDs are errors, not overwritten."""
    zip_inventory(content)
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        sheet = workbook[sheet_name]
        if (sheet.max_row or 0) * (sheet.max_column or 0) > 1_000_000:
            raise ValueError("Překročen limit buněk.")
        rows = sheet.iter_rows(min_row=header_row)
        headers = next(rows)
        labels = {i: str(c.value).strip() for i, c in enumerate(headers) if c.value is not None}
        if any(not v for v in labels.values()) or len(set(labels.values())) != len(labels):
            raise ValueError("Nejednoznačné hlavičky technické tabulky.")
        result = {}
        for number, row in enumerate(rows, header_row + 1):
            key = row[key_column - 1].value
            if key is None or not str(key).strip():
                continue
            key = str(key).strip()
            if key in result:
                raise ValueError(f"Duplicitní ID v listu {sheet_name}: {key}")
            result[key] = {"sheet": sheet_name, "row": number,
                           "fields": {labels[i]: {"cell": f"{get_column_letter(i + 1)}{number}", "value": cell_value(row[i].value),
                                                   "formula": row[i].data_type == "f"} for i in labels}}
        return result
    finally:
        workbook.close()


def record_changes(before: dict, after: dict) -> dict:
    """ID based diff; compare values, not shifted row coordinates."""
    changed = {}
    for key in sorted(before.keys() & after.keys()):
        left = {k: v["value"] for k, v in before[key]["fields"].items()}
        right = {k: v["value"] for k, v in after[key]["fields"].items()}
        fields = {f: {"before": left.get(f), "after": right.get(f)} for f in left.keys() | right.keys() if left.get(f) != right.get(f)}
        if fields:
            changed[key] = fields
    return {"added": sorted(after.keys() - before.keys()), "removed": sorted(before.keys() - after.keys()), "changed": changed}
