"""Import bez interpretace osobních identifikátorů jako čísel."""
import csv
import io
from pathlib import Path

from app.models import Dataset, Row

MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100_000


def _headers(values: list[str]) -> tuple[str, ...]:
    names = tuple(v.strip() for v in values)
    if not names or any(not n for n in names) or len(set(names)) != len(names):
        raise ValueError("Hlavičky musí být neprázdné a jedinečné.")
    return names


def read_csv(path: Path, encoding: str = "auto", delimiter: str = "auto") -> Dataset:
    if path.stat().st_size > MAX_BYTES:
        raise ValueError("Soubor překračuje limit 25 MiB.")
    raw = path.read_bytes()
    if encoding == "auto":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1250")
    elif encoding in {"utf-8-sig", "cp1250"}:
        text = raw.decode(encoding)
    else:
        raise ValueError("Nepodporované kódování.")
    if delimiter == "auto":
        try:
            delimiter = csv.Sniffer().sniff(text[:65536], delimiters=";,\t|").delimiter
        except csv.Error as exc:
            raise ValueError("Nelze určit oddělovač. Vyberte jej ručně.") from exc
    if delimiter not in {";", ",", "\t", "|"}:
        raise ValueError("Nepodporovaný oddělovač.")
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    try:
        names = _headers(next(reader))
        rows = []
        while True:
            number = reader.line_num + 1
            values = next(reader, None)
            if values is None:
                break
            if not values:
                continue
            if len(values) != len(names):
                raise ValueError(f"Řádek {number}: počet buněk neodpovídá hlavičce.")
            rows.append(Row(number, dict(zip(names, values))))
            if len(rows) > MAX_ROWS:
                raise ValueError("Překročen limit 100 000 záznamů.")
    except StopIteration as exc:
        raise ValueError("Prázdný soubor.") from exc
    except csv.Error as exc:
        raise ValueError(f"Neplatné CSV u řádku {reader.line_num}.") from exc
    if not rows:
        raise ValueError("Soubor neobsahuje datové řádky.")
    return Dataset(names, tuple(rows))


def read_xlsx(path: Path) -> Dataset:
    from openpyxl import load_workbook
    from zipfile import ZipFile

    if path.stat().st_size > MAX_BYTES:
        raise ValueError("Soubor překračuje limit 25 MiB.")
    with ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("Rozbalený XLSX překračuje limit 100 MiB.")
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheet = workbook.active
        if sheet.max_row and sheet.max_row > MAX_ROWS + 1:
            raise ValueError("Překročen limit 100 000 záznamů.")
        cells = sheet.iter_rows()
        header = next(cells)
        if any(c.value is not None and c.data_type != "s" for c in header):
            raise ValueError("Hlavičky XLSX musí být textové.")
        names = _headers([c.value or "" for c in header])
        rows = []
        for number, row in enumerate(cells, 2):
            if all(c.value is None for c in row):
                continue
            values, numeric = {}, set()
            for name, cell in zip(names, row):
                if cell.value is None:
                    values[name] = ""
                elif cell.data_type == "s":
                    values[name] = cell.value
                elif cell.data_type == "n" and isinstance(cell.value, (int, float)):
                    from decimal import Decimal
                    value = Decimal(str(cell.value))
                    if not value.is_finite():
                        raise ValueError(f"Řádek {number}: nekonečná číselná hodnota.")
                    values[name] = format(value, "f")
                    numeric.add(name)
                else:
                    raise ValueError(f"Řádek {number}: vzorce, chyby, logické a datumové buňky nejsou podporovány.")
            rows.append(Row(number, values, frozenset(numeric)))
            if len(rows) > MAX_ROWS:
                raise ValueError("Překročen limit záznamů.")
        if not rows:
            raise ValueError("Soubor neobsahuje datové řádky.")
        return Dataset(names, tuple(rows))
    except StopIteration as exc:
        raise ValueError("Prázdný sešit.") from exc
    finally:
        workbook.close()
